import pandas as pd
from openpyxl import load_workbook
from engine.python.export_xlsx import export_workbook


def test_gene_names_match_case_insensitively(tmp_path):
    res = tmp_path / "results"
    res.mkdir()
    pd.DataFrame({"Gene": ["mab0001", "MAB0002c"], "log2FoldChange": [1.0, -1.0],
                  "padj": [0.01, 0.2]}).to_csv(res / "B_vs_A.tsv", sep="\t", index=False)
    out = tmp_path / "r.xlsx"
    export_workbook(str(res), str(out), gene_names={"MAB0001": "dnaA", "MAB0002c": "dnaN"})
    ws = load_workbook(out)["B_vs_A"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert [r[1] for r in rows] == ["dnaA", "dnaN"]
