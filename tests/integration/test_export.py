import glob
import pytest
from pathlib import Path
from openpyxl import load_workbook
from engine.python.gene_names import gene_names_from_gff
from engine.python.export_xlsx import export_workbook

RES = "/path/to/projects/rnaseq_mabs_media/06_deseq/results"
GFF = "/path/to/projects/bac-rnaseq/refs/mabs/NC_010397.1.gff3"
pytestmark = pytest.mark.skipif(not (glob.glob(RES + "/*.tsv") and Path(GFF).exists()),
                                reason="need real results + bundled GFF")


def test_gene_names_and_workbook(tmp_path):
    gn = gene_names_from_gff(GFF)
    assert len(gn) > 1000
    out = tmp_path / "results.xlsx"
    export_workbook(RES, str(out), gene_names=gn)
    wb = load_workbook(out)
    assert "summary" in wb.sheetnames
    contrast_sheets = [s for s in wb.sheetnames if s != "summary"]
    assert contrast_sheets
    header = [c.value for c in wb[contrast_sheets[0]][1]]
    assert "gene_name" in header and "Gene" in header
